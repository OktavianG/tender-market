"""Исследование рынка пластмассовых изделий по данным ЕИС и Tender.Pro.

Зависимости:
    pip install requests beautifulsoup4 openpyxl

На первом запуске обрабатываются только две страницы каждого поискового
запроса.  Это намеренное тестовое ограничение (MAX_PAGES_PER_QUERY).
"""

from __future__ import annotations

import re
import hashlib
import random
import time
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlencode, urljoin, urlparse, urlunparse

import requests
import urllib3
from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from sources.tender_pro import collect_tender_pro


# ---------------------------------------------------------------------------
# Настройки
# ---------------------------------------------------------------------------

BASE_URL = "https://zakupki.gov.ru"
SEARCH_URL = f"{BASE_URL}/epz/order/extendedsearch/results.html"
DATE_FROM = "01.01.2024"
DATE_TO = date.today().strftime("%d.%m.%Y")
MAX_PAGES_PER_QUERY = 2
# ЕИС блокирует адреса при частых автоматических обращениях. Для безопасного
# теста обрабатываем не более 30 карточек и делаем длинные случайные паузы.
MAX_EIS_TENDERS = 30
EIS_DELAY_MIN = 5.0
EIS_DELAY_MAX = 8.0
TENDER_PRO_DELAY = 1.2
TIMEOUT = 40
OUTPUT_FILE = Path(__file__).with_name("plastics_tenders.xlsx")
CACHE_DIR = Path(__file__).with_name("cache_eis")
ENABLE_TENDER_PRO = True

# Коды субъектов в параметре customerPlaceCodes ЕИС.
REGIONS = {
    "Республика Татарстан": "92000000000",
    "Самарская область": "36000000000",
}

SEARCH_QUERIES = [
    "пластмассовый",
    "пластиковый",
    "полипропиленовый",
    "полиэтиленовый",
    "заглушка пластиковая",
    "втулка пластиковая",
    "ящик пластмассовый",
    "лоток пластиковый",
    "корпус пластиковый",
    "контейнер пластиковый",
    "крепеж пластиковый",
    "детали из пластмасс",
]

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36"
    ),
    "Accept-Language": "ru-RU,ru;q=0.9",
}

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


def make_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(HEADERS)
    retry = Retry(
        total=4,
        connect=4,
        read=4,
        status=4,
        backoff_factor=1.0,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset(("GET",)),
        raise_on_status=False,
    )
    session.mount("https://", HTTPAdapter(max_retries=retry))
    return session


SESSION = make_session()


class EISBlockedError(RuntimeError):
    """ЕИС вернул страницу блокировки вместо запрошенных данных."""


def eis_pause() -> None:
    time.sleep(random.uniform(EIS_DELAY_MIN, EIS_DELAY_MAX))


# ---------------------------------------------------------------------------
# Общие утилиты
# ---------------------------------------------------------------------------

def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip()


def parse_number(value: Any) -> float | None:
    text = clean_text(value)
    if not text or text in {"-", "—"}:
        return None
    text = re.sub(r"[^\d,.-]", "", text).replace(",", ".")
    # Если разделителей несколько, последний считаем десятичным.
    if text.count(".") > 1:
        parts = text.split(".")
        text = "".join(parts[:-1]) + "." + parts[-1]
    try:
        return float(text)
    except ValueError:
        return None


def extract_codes(value: str) -> tuple[str, str]:
    value = clean_text(value)
    ktru_match = re.search(r"\b\d{2}\.\d{2}(?:\.\d{2})?(?:\.\d{3})?-\d+\b", value)
    okpd_match = re.search(r"\b\d{2}\.\d{2}(?:\.\d{2})?(?:\.\d{3})?\b", value)
    return (okpd_match.group(0) if okpd_match else "", ktru_match.group(0) if ktru_match else "")


def get_html(url: str, params: dict[str, Any] | None = None) -> str | None:
    prepared_url = requests.Request("GET", url, params=params).prepare().url or url
    cache_key = hashlib.sha256(prepared_url.encode("utf-8")).hexdigest()
    cache_file = CACHE_DIR / f"{cache_key}.html"
    if cache_file.exists():
        print(f"КЭШ: {prepared_url}")
        return cache_file.read_text(encoding="utf-8")
    try:
        response = SESSION.get(url, params=params, timeout=TIMEOUT, verify=False)
        print(f"HTTP {response.status_code}: {response.url}")
        response.raise_for_status()
        response.encoding = response.apparent_encoding or response.encoding
        html = response.text
        lowered = html.lower()
        if "custom ip address blocked" in lowered or "был заблокирован в связи с подозрительной активностью" in lowered:
            raise EISBlockedError(
                "ЕИС заблокировал IP. Сбор остановлен; повторный запуск сейчас запрещён."
            )
        CACHE_DIR.mkdir(exist_ok=True)
        cache_file.write_text(html, encoding="utf-8")
        return html
    except requests.RequestException as exc:
        print(f"ОШИБКА HTTP: {exc}")
        return None


def canonical_common_info_url(href: str, reg_number: str) -> str:
    """Сохраняет рабочий путь карточки, меняя только вкладку на common-info."""
    full = urljoin(BASE_URL, href)
    parsed = urlparse(full)
    path = parsed.path
    if "/notice/" in path and "/view/" in path:
        path = re.sub(r"/view/[^/]+\.html$", "/view/common-info.html", path)
    query = parse_qs(parsed.query)
    query["regNumber"] = [reg_number]
    return urlunparse(parsed._replace(path=path, query=urlencode(query, doseq=True), fragment=""))


# ---------------------------------------------------------------------------
# Поиск закупок
# ---------------------------------------------------------------------------

def search_tenders() -> dict[str, dict[str, Any]]:
    tenders: dict[str, dict[str, Any]] = {}
    total_runs = len(SEARCH_QUERIES) * len(REGIONS)
    run_no = 0

    for region_name, region_code in REGIONS.items():
        for query in SEARCH_QUERIES:
            run_no += 1
            before_query = len(tenders)
            print(f"\n[{run_no}/{total_runs}] {region_name}: {query}")

            for page in range(1, MAX_PAGES_PER_QUERY + 1):
                params = {
                    "fz44": "on",
                    "fz223": "on",
                    "searchString": query,
                    "publishDateFrom": DATE_FROM,
                    "publishDateTo": DATE_TO,
                    "customerPlaceWithNested": "on",
                    "customerPlaceCodes": region_code,
                    "pageNumber": page,
                    "recordsPerPage": "_10",
                }
                html = get_html(SEARCH_URL, params=params)
                if not html:
                    continue
                soup = BeautifulSoup(html, "html.parser")
                found_on_page = 0

                for link in soup.select('a[href*="regNumber="]'):
                    href = link.get("href", "")
                    if "/notice/" not in href:
                        continue
                    # На странице поиска есть несколько ссылок с regNumber:
                    # карточка, печатная форма, документы и служебные окна.
                    # Товарные позиции находятся только на common-info.html.
                    # Раньше первой попадалась printForm/listModal.html, поэтому
                    # все 169 закупок открывались без товарной таблицы.
                    if "/view/common-info.html" not in href:
                        continue
                    match = re.search(r"(?:[?&])regNumber=(\d+)", href)
                    if not match:
                        continue
                    reg_number = match.group(1)
                    if reg_number not in tenders:
                        found_on_page += 1
                        tenders[reg_number] = {
                            "reg_number": reg_number,
                            "url": canonical_common_info_url(href, reg_number),
                            "queries": set(),
                            "regions": set(),
                        }
                    tenders[reg_number]["queries"].add(query)
                    tenders[reg_number]["regions"].add(region_name)

                print(f"  страница {page}: новых закупок {found_on_page}")
                eis_pause()

            print(f"  новых по запросу после дедупликации: {len(tenders) - before_query}")

    return tenders


# ---------------------------------------------------------------------------
# Разбор common-info.html
# ---------------------------------------------------------------------------

def find_value_by_label(soup: BeautifulSoup, labels: tuple[str, ...]) -> str:
    """Извлекает реквизит независимо от div/table-вёрстки ЕИС."""
    labels_lower = tuple(label.lower() for label in labels)
    for node in soup.find_all(["div", "span", "td", "th", "dt"]):
        own = clean_text(node.get_text(" ", strip=True)).lower()
        if not any(own == label or own.startswith(label + ":") for label in labels_lower):
            continue
        # Частая вёрстка: label и value — соседние блоки/ячейки.
        sibling = node.find_next_sibling()
        if sibling:
            value = clean_text(sibling.get_text(" ", strip=True))
            if value and value.lower() not in labels_lower:
                return value
        parent = node.parent
        if parent:
            whole = clean_text(parent.get_text(" ", strip=True))
            label_text = clean_text(node.get_text(" ", strip=True))
            value = clean_text(whole.removeprefix(label_text).lstrip(": "))
            if value:
                return value
    return ""


def find_product_table(soup: BeautifulSoup) -> Tag | None:
    best: tuple[int, Tag | None] = (0, None)
    markers = (
        "код позиции", "наименование товара", "наименование товара, работы, услуги",
        "ед. измерения", "единица измерения", "количество", "цена за ед", "стоимость",
    )
    for table in soup.find_all("table"):
        text = clean_text(table.get_text(" ", strip=True)).lower()
        score = sum(marker in text for marker in markers)
        if score > best[0]:
            best = (score, table)
    return best[1] if best[0] >= 4 else None


def header_key(text: str) -> str | None:
    text = clean_text(text).lower()
    if "код позиции" in text or "окпд" in text or "ктру" in text:
        return "code"
    if "наименование" in text and any(x in text for x in ("товар", "работ", "услуг")):
        return "name"
    if "ед." in text or "единица измерения" in text:
        return "unit"
    if "количество" in text or "объем" in text:
        return "quantity"
    if "цена" in text and ("ед" in text or "единиц" in text):
        return "price"
    if "стоимость" in text or "сумма" in text:
        return "amount"
    return None


def detect_columns(table: Tag) -> dict[str, int]:
    best: dict[str, int] = {}
    for row in table.find_all("tr")[:8]:
        cells = row.find_all(["th", "td"], recursive=False)
        current = {key: i for i, cell in enumerate(cells) if (key := header_key(cell.get_text(" ", strip=True)))}
        if len(current) > len(best):
            best = current
    return best


def parse_product_table(table: Tag, tender: dict[str, Any], meta: dict[str, str]) -> list[dict[str, Any]]:
    products: list[dict[str, Any]] = []
    columns = detect_columns(table)

    for tr in table.find_all("tr"):
        cells = tr.find_all(["td", "th"], recursive=False)
        values = [clean_text(cell.get_text(" ", strip=True)) for cell in cells]
        if len(values) < 5 or all(not value for value in values):
            continue

        if {"code", "name"}.issubset(columns) and max(columns.values()) < len(values):
            get = lambda key: values[columns[key]] if key in columns and columns[key] < len(values) else ""
            code_raw, name = get("code"), get("name")
            unit, quantity_raw = get("unit"), get("quantity")
            price_raw, amount_raw = get("price"), get("amount")
        else:
            # Резервный режим — прежняя рабочая структура таблицы ЕИС.
            offset = 1 if values and not values[0] else 0
            if len(values) < offset + 6:
                continue
            code_raw, name, unit, quantity_raw, price_raw, amount_raw = values[offset:offset + 6]

        if not re.search(r"\b\d{2}\.\d{2}", code_raw) or not name:
            continue
        if any(x in name.lower() for x in ("наименование характеристики", "инструкция по заполнению")):
            continue

        okpd, ktru = extract_codes(code_raw)
        quantity = parse_number(quantity_raw)
        price = parse_number(price_raw)
        amount = parse_number(amount_raw)
        if amount is None and quantity is not None and price is not None:
            amount = quantity * price

        product = {
            "source": "ЕИС",
            "reg_number": tender["reg_number"],
            "law": meta.get("law", ""),
            "customer": meta.get("customer", ""),
            "region": ", ".join(sorted(tender["regions"])),
            "search_queries": ", ".join(sorted(tender["queries"])),
            "okpd": okpd,
            "ktru": ktru,
            "name": name,
            "unit": unit,
            "quantity": quantity,
            "price": price,
            "amount": amount,
            "url": tender["url"],
        }
        product["category"], product["reason"] = classify_product(product)
        products.append(product)
    return products


def parse_tender(tender: dict[str, Any]) -> list[dict[str, Any]]:
    html = get_html(tender["url"])
    if not html:
        return []
    soup = BeautifulSoup(html, "html.parser")
    law_text = find_value_by_label(soup, ("Закон", "Размещение осуществляет"))
    page_text = clean_text(soup.get_text(" ", strip=True))
    if "44-ФЗ" in law_text or "44-ФЗ" in page_text:
        law = "44-ФЗ"
    elif "223-ФЗ" in law_text or "223-ФЗ" in page_text:
        law = "223-ФЗ"
    else:
        law = ""
    customer = find_value_by_label(
        soup,
        ("Наименование организации", "Заказчик", "Организация, осуществляющая размещение"),
    )
    table = find_product_table(soup)
    if table is None:
        print("  таблица товарных позиций не найдена")
        return []
    products = parse_product_table(table, tender, {"law": law, "customer": customer})
    print(f"  товарных позиций: {len(products)}")
    return products


# ---------------------------------------------------------------------------
# Классификация
# ---------------------------------------------------------------------------

EXCLUDE_GROUPS = {
    "медицина/стоматология": (
        "медицинск", "стоматолог", "зуботех", "лабораторн", "шприц", "катетер",
        "канюл", "пробирк", "наконечник пипет", "биоматериал", "дезинфек",
    ),
    "пищевые товары": (
        "продукт питан", "напиток", "молоко", "сок ", "крупа", "макарон", "консервы",
        "хлеб", "масло раститель", "мясо", "рыба", "овощ", "фрукт", "кондитер",
    ),
    "бытовая мелочёвка": (
        "одноразов", "стакан", "тарелк", "ложка", "вилка", "ведро", "таз ",
        "корзина для мусора", "щетка", "расческ", "игрушк", "канцеляр",
    ),
    "плёнка/упаковка": (
        "пленк", "плёнк", "стрейч", "пакет", "мешок", "упаковоч", "рукав полимер",
    ),
    "трубы/профили": ("труб", "шланг", "профил", "гофр", "фитинг", "муфт"),
    "кабельная продукция": ("кабель", "провод", "оптоволок", "изоляция провод"),
    "работы/услуги": (
        "услуг", "работы по", "ремонт", "монтаж", "демонтаж", "обслуживан", "аренд",
        "утилизац", "изготовление по техническому заданию",
    ),
    "сырьё/материалы": (
        "гранул", "смола", "компаунд", "сырье", "сырьё", "порошок", "клей", "герметик",
        "лак ", "краска", "паста", "силикон", "полимерный материал", "листовой пластик",
    ),
}

TARGET_WORDS = (
    "заглушк", "втулк", "корпус", "крышк", "ящик", "лоток", "контейнер",
    "крепеж", "крепёж", "кронштейн", "держатель", "фиксатор", "зажим", "ручка",
    "опора", "колпачок", "кожух", "кассета", "поддон", "деталь", "комплектующ",
    "вставка", "переходник", "шайба", "проставка", "клипса", "уголок",
)
PLASTIC_WORDS = (
    "пластик", "пластмасс", "полимер", "полипропилен", "полиэтилен", "полиамид",
    "абс", "abs", "пвх", "пнд", "пэт", "термопласт",
)


def classify_product(product: dict[str, Any]) -> tuple[str, str]:
    text = clean_text(f"{product.get('name', '')} {product.get('okpd', '')}").lower()
    for reason, words in EXCLUDE_GROUPS.items():
        if any(word in text for word in words):
            return "ИСКЛЮЧИТЬ", reason
    # Разделы 21 и 32.5 преимущественно медицинские; это страховка от новых терминов.
    if product.get("okpd", "").startswith(("21.", "32.5")):
        return "ИСКЛЮЧИТЬ", "медицинский ОКПД2"
    has_target = any(word in text for word in TARGET_WORDS)
    has_plastic = any(word in text for word in PLASTIC_WORDS) or product.get("okpd", "").startswith("22.29")
    if has_target and has_plastic:
        return "ЦЕЛЕВОЕ", "техническое формованное изделие из пластмассы"
    if has_plastic:
        return "ПРОВЕРИТЬ", "пластмассовое изделие без однозначного типа"
    if has_target:
        return "ПРОВЕРИТЬ", "возможное техническое изделие; материал не подтверждён"
    return "ИСКЛЮЧИТЬ", "нет признаков целевого формованного изделия"


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

COLUMNS = [
    ("source", "Источник"), ("category", "Категория"), ("reason", "Причина классификации"),
    ("reg_number", "№ закупки"), ("law", "Закон"), ("customer", "Заказчик"),
    ("region", "Регион поиска"), ("okpd", "ОКПД2"), ("ktru", "КТРУ"),
    ("name", "Наименование / описание"), ("unit", "Ед. измерения"),
    ("quantity", "Количество"), ("price", "Цена за единицу, ₽"),
    ("amount", "Сумма, ₽"), ("search_queries", "Найдено по запросам"), ("url", "Ссылка"),
]


def normalized_product_name(value: str) -> str:
    value = clean_text(value).lower().replace("ё", "е")
    value = re.sub(r"[^а-яa-z0-9]+", " ", value)
    return clean_text(value)[:250]


def style_sheet(ws, widths: dict[int, int] | None = None) -> None:
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = (widths or {}).get(col, 18)


def add_products_sheet(wb: Workbook, title: str, products: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append([label for _, label in COLUMNS])
    for product in products:
        ws.append([product.get(key) for key, _ in COLUMNS])
        link = ws.cell(ws.max_row, len(COLUMNS))
        link.value = "Открыть закупку"
        link.hyperlink = product["url"]
        link.style = "Hyperlink"
    for row in range(2, ws.max_row + 1):
        for col in (12, 13, 14):
            ws.cell(row, col).number_format = '#,##0.00'
    style_sheet(ws, {1: 16, 3: 34, 4: 23, 6: 45, 10: 70, 15: 42, 16: 22})


def add_product_summary(wb: Workbook, products: list[dict[str, Any]]) -> None:
    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    for p in products:
        key = (normalized_product_name(p["name"]), p["okpd"])
        item = grouped.setdefault(key, {"name": p["name"], "okpd": p["okpd"], "tenders": set(), "qty": 0.0, "amount": 0.0})
        item["tenders"].add(p["reg_number"])
        item["qty"] += p["quantity"] or 0
        item["amount"] += p["amount"] or 0
    ws = wb.create_sheet("Сводка по товарам")
    ws.append(["Наименование товара", "ОКПД2", "Закупок", "Количество", "Сумма, ₽"])
    for item in sorted(grouped.values(), key=lambda x: x["amount"], reverse=True):
        ws.append([item["name"], item["okpd"], len(item["tenders"]), item["qty"], item["amount"]])
    style_sheet(ws, {1: 75, 2: 18, 3: 14, 4: 18, 5: 22})


def add_okpd_summary(wb: Workbook, products: list[dict[str, Any]]) -> None:
    grouped: dict[str, dict[str, Any]] = defaultdict(lambda: {"positions": 0, "tenders": set(), "amount": 0.0})
    for p in products:
        item = grouped[p["okpd"] or "Без кода"]
        item["positions"] += 1
        item["tenders"].add(p["reg_number"])
        item["amount"] += p["amount"] or 0
    ws = wb.create_sheet("Сводка ОКПД2")
    ws.append(["ОКПД2", "Позиций", "Закупок", "Сумма, ₽"])
    for code, item in sorted(grouped.items(), key=lambda x: x[1]["amount"], reverse=True):
        ws.append([code, item["positions"], len(item["tenders"]), item["amount"]])
    style_sheet(ws, {1: 20, 2: 14, 3: 14, 4: 22})


def add_customer_summary(wb: Workbook, products: list[dict[str, Any]]) -> None:
    grouped: dict[str, dict[str, Any]] = defaultdict(lambda: {"positions": 0, "tenders": set(), "amount": 0.0})
    for p in products:
        customer = p["customer"] or "Не удалось определить"
        item = grouped[customer]
        item["positions"] += 1
        item["tenders"].add(p["reg_number"])
        item["amount"] += p["amount"] or 0
    ws = wb.create_sheet("Заказчики")
    ws.append(["Заказчик", "Закупок", "Позиций", "Сумма, ₽"])
    for customer, item in sorted(grouped.items(), key=lambda x: x[1]["amount"], reverse=True):
        ws.append([customer, len(item["tenders"]), item["positions"], item["amount"]])
    style_sheet(ws, {1: 80, 2: 14, 3: 14, 4: 22})


def save_excel(products: list[dict[str, Any]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    add_products_sheet(wb, "Целевые изделия", [p for p in products if p["category"] == "ЦЕЛЕВОЕ"])
    add_products_sheet(wb, "Проверить", [p for p in products if p["category"] == "ПРОВЕРИТЬ"])
    add_products_sheet(wb, "Все позиции", products)
    relevant = [p for p in products if p["category"] != "ИСКЛЮЧИТЬ"]
    add_product_summary(wb, relevant)
    add_okpd_summary(wb, relevant)
    add_customer_summary(wb, relevant)
    wb.save(OUTPUT_FILE)


def main() -> None:
    print("=" * 80)
    print("ПОИСК РЫНКА ПЛАСТМАССОВЫХ ИЗДЕЛИЙ")
    print("=" * 80)
    print("Регионы:", ", ".join(REGIONS))
    print("Законы: 44-ФЗ, 223-ФЗ")
    print(f"Период: {DATE_FROM} — {DATE_TO}")
    print(f"Тестовый лимит: {MAX_PAGES_PER_QUERY} страницы на запрос и регион")

    tenders = search_tenders()
    print(f"\nУНИКАЛЬНЫХ ЗАКУПОК: {len(tenders)}")

    products: list[dict[str, Any]] = []
    eis_tenders = list(tenders.values())[:MAX_EIS_TENDERS]
    print(f"БЕЗОПАСНЫЙ ТЕСТ: обрабатываем {len(eis_tenders)} из {len(tenders)} карточек ЕИС")
    for index, tender in enumerate(eis_tenders, 1):
        print(f"\n[{index}/{len(eis_tenders)}] {tender['reg_number']}")
        products.extend(parse_tender(tender))
        eis_pause()

    if ENABLE_TENDER_PRO:
        products.extend(
            collect_tender_pro(
                session=SESSION,
                search_queries=SEARCH_QUERIES,
                classify=classify_product,
                max_pages=MAX_PAGES_PER_QUERY,
                delay=TENDER_PRO_DELAY,
                timeout=TIMEOUT,
            )
        )

    save_excel(products)
    counts = Counter(p["category"] for p in products)
    print("\n" + "=" * 80)
    print(f"ПОЗИЦИЙ ВСЕГО: {len(products)}")
    print(f"ЦЕЛЕВОЕ: {counts['ЦЕЛЕВОЕ']}")
    print(f"ПРОВЕРИТЬ: {counts['ПРОВЕРИТЬ']}")
    print(f"ИСКЛЮЧИТЬ: {counts['ИСКЛЮЧИТЬ']}")
    print(f"EXCEL СОХРАНЁН: {OUTPUT_FILE}")


if __name__ == "__main__":
    try:
        main()
    except EISBlockedError as exc:
        print("\n" + "=" * 80)
        print("СБОР БЕЗОПАСНО ОСТАНОВЛЕН")
        print(exc)
        print("Не запускайте программу повторно до снятия блокировки ЕИС.")
